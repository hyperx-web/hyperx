# Insert the recorded research animations (MP4) into slides 7-14 of the lab introduction deck, large, auto-playing in a loop.
# Usage: python tools/insert_videos.py <in.pptx> <out.pptx> <mp4_dir> <poster_dir>
import os, sys, copy
from lxml import etree
from pptx import Presentation
from pptx.util import Inches

SLIDES = {7: "ai-rlbf", 8: "ai-btnws", 9: "ai-agent", 10: "6g-rsma", 11: "6g-afdm",
          12: "sat-constellation", 13: "sat-beams", 14: "sat-threebody"}
AREA = (0.45, 1.75, 12.43, 5.5)            # x, y, w, h (inches) below the section label
NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
P = "{%s}" % NS["p"]


def autoplay_loop(slide, shape_id):
    """Replace the slide timing with: play the movie automatically, repeat until the slide ends."""
    sld = slide._element
    old = sld.find(P + "timing")
    if old is not None:
        sld.remove(old)
    xml = f'''<p:timing xmlns:p="{NS['p']}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
  <p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
    <p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>
      <p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
        <p:par><p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
          <p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="6" dur="1" fill="hold"/><p:tgtEl><p:spTgt spid="{shape_id}"/></p:tgtEl></p:cBhvr></p:cmd>
        </p:childTnLst></p:cTn></p:par>
      </p:childTnLst></p:cTn></p:par>
    </p:childTnLst></p:cTn></p:par>
  </p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>
  <p:video><p:cMediaNode vol="80000" mute="1"><p:cTn id="7" repeatCount="indefinite" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{shape_id}"/></p:tgtEl></p:cMediaNode></p:video>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'''
    t = etree.fromstring(xml)
    ext = sld.find(P + "extLst")
    if ext is not None:
        ext.addprevious(t)
    else:
        sld.append(t)


def main():
    src, dst, mp4dir, posterdir = sys.argv[1:5]
    prs = Presentation(src)
    for num, name in SLIDES.items():
        slide = prs.slides[num - 1]
        poster = os.path.join(posterdir, name + ".png")
        from PIL import Image
        w_px, h_px = Image.open(poster).size
        x, y, w, h = AREA
        if w_px / h_px > w / h:
            vw, vh = w, w * h_px / w_px
        else:
            vw, vh = h * w_px / h_px, h
        left, top = x + (w - vw) / 2, y + (h - vh) / 2
        mv = slide.shapes.add_movie(os.path.join(mp4dir, name + ".mp4"), Inches(left), Inches(top), Inches(vw), Inches(vh),
                                    poster_frame_image=poster, mime_type="video/mp4")
        mv.name = "Animation " + name
        autoplay_loop(slide, mv.shape_id)
        print(num, name, f"{vw:.2f}x{vh:.2f} in")
    prs.save(dst)


if __name__ == "__main__":
    main()
